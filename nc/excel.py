"""점검 회차 1건 → 엑셀 보고서(.xlsx).

시트
- 요약: 장비별 상태, CPU·메모리, 이상 항목
- 백본 CPU·메모리: 슬롯별 Last 5 sec / 1 min / 5 min, Total / Used / Free / FreeRatio
- 컨트롤러 CPU·메모리: MM·WLC별 CPU user / system / idle, Memory Total / Used / Free / Used%
- 백본 로그: show logbuffer reverse (최근순, 오류급 이상 강조)
- 백본 VRRP: 그룹별 현재 역할, 장비 로그의 VRRP 상태 변경 시각
- 점검 항목: 모든 판단 근거
- 원본 출력: 장비가 보낸 실제 응답(확인용)
"""

from __future__ import annotations

import io
import re
import tempfile
import time
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .evaluate import LEVEL_LABEL, fmt_uptime

FONT = "맑은 고딕"
INK = "172033"
FILL = {"ok": "DCEFE6", "info": "DCEFE6", "warn": "F6E8CC", "crit": "F8DCD8", "unknown": "E6E1F3"}
TEXT = {"ok": "1B8460", "info": "1B8460", "warn": "8A5600", "crit": "B42318", "unknown": "5B4A99"}
ROLE_ORDER = {"BB": 0, "MM": 1, "WLC": 2}
ROLE_NAME = {"BB": "백본", "MM": "MM", "WLC": "WLC"}
THIN = Side(style="thin", color="CCD3DC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PCT = "0.0%"
NUM = "#,##0"
ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _clean(v: Any) -> Any:
    if isinstance(v, str):
        v = ILLEGAL.sub("", v)
        return v[:32000]
    return v


def _pct(v: Any) -> float | None:
    return None if v is None else round(float(v) / 100, 4)


def _mb(kb: Any) -> float | None:
    return None if kb is None else round(int(kb) / 1024, 1)


class Sheet:
    def __init__(self, ws, title: str, headers: list[tuple[str, int]], subtitle: str = "") -> None:
        self.ws = ws
        ws.sheet_view.showGridLines = False
        ws["A1"] = title
        ws["A1"].data_type = "s"
        ws["A1"].font = Font(name=FONT, size=14, bold=True, color=INK)
        if subtitle:
            ws["A2"] = subtitle
            ws["A2"].data_type = "s"
            ws["A2"].font = Font(name=FONT, size=9, color="5D6879")
        self.hrow = 4
        for i, (h, w) in enumerate(headers, start=1):
            c = ws.cell(row=self.hrow, column=i, value=h)
            c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=INK)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BORDER
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[self.hrow].height = 30
        ws.freeze_panes = ws.cell(row=self.hrow + 1, column=1)
        self.row = self.hrow
        self.ncol = len(headers)

    def add(self, values: list[Any], fmts: dict[int, str] | None = None, level: str | None = None,
            level_col: int | None = None, wrap: set[int] | None = None, row_fill: str | None = None) -> None:
        self.row += 1
        for i, v in enumerate(values, start=1):
            value = _clean(v)
            c = self.ws.cell(row=self.row, column=i, value=value)
            # Device names, CLI output, and logs are data, never Excel formulas.
            # Explicit string typing preserves the original text without a visible prefix.
            if isinstance(value, str):
                c.data_type = "s"
            c.font = Font(name=FONT, size=10)
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=bool(wrap and i in wrap))
            if fmts and i in fmts and v is not None:
                c.number_format = fmts[i]
                c.alignment = Alignment(horizontal="right", vertical="top")
            if row_fill:
                c.fill = PatternFill("solid", fgColor=row_fill)
        if level and level_col:
            c = self.ws.cell(row=self.row, column=level_col)
            c.fill = PatternFill("solid", fgColor=FILL.get(level, "FFFFFF"))
            c.font = Font(name=FONT, size=10, bold=True, color=TEXT.get(level, INK))
            c.alignment = Alignment(horizontal="center", vertical="top")

    def finish(self, note: str = "") -> None:
        if self.row > self.hrow:
            self.ws.auto_filter.ref = f"A{self.hrow}:{get_column_letter(self.ncol)}{self.row}"
        if note:
            r = self.row + 2
            self.ws.cell(row=r, column=1, value=note).font = Font(name=FONT, size=9, color="5D6879")


def _kind(d: dict) -> str:
    """구분: 백본 / MM / WLC IN / WLC OUT …"""
    if d["role"] == "WLC" and d.get("cluster"):
        return f"WLC {d['cluster']}"
    return ROLE_NAME.get(d["role"], d["role"])


def _level_for(value: float | None, warn: float, crit: float) -> str | None:
    if value is None:
        return None
    return "crit" if value >= crit else "warn" if value >= warn else "ok"


def build_workbook(run: dict, site: str, thresholds: dict, raw: Callable[[str, str], dict | None]) -> bytes:
    started = time.strftime("%Y-%m-%d %H:%M", time.localtime(run["started"]))
    s = run["summary"]
    devs = sorted(run["devices"].values(), key=lambda d: (ROLE_ORDER.get(d["role"], 9), d.get("cluster", ""), d["name"]))
    th = thresholds
    sub = f"{site}  |  점검 시각 {started}  |  회차 #{run['id']}  |  {'자동' if run.get('trigger') == 'auto' else '수동'} 점검"

    wb = Workbook()

    # ------------------------------------------------------------------ 요약
    ws = wb.active
    ws.title = "요약"
    sh = Sheet(ws, f"{site} 무선망 점검 보고서", [
        ("구분", 10), ("장비", 16), ("IP", 15), ("상태", 10), ("CPU (%)", 10), ("메모리 사용 (%)", 12),
        ("가동시간", 13), ("이상 항목 수", 11), ("확인할 내용", 70)], sub)
    c = s.get("counts", {})
    ws["H1"] = "종합"
    ws["H1"].font = Font(name=FONT, size=10, color="5D6879")
    ws["I1"] = (f"{LEVEL_LABEL[s['overall']]}  (장애 {c.get('crit', 0)} · 주의 {c.get('warn', 0)} · "
                f"확인 불가 {c.get('unknown', 0)} · 정상 {c.get('ok', 0)})"
                + (f"  |  Active Client {s['total_active']:,}명" if s.get("total_active") is not None else ""))
    ws["I1"].font = Font(name=FONT, size=11, bold=True, color=TEXT.get(s["overall"], INK))
    for d in devs:
        m = d.get("metrics", {})
        bad = [ch for ch in d["checks"] if ch["level"] in ("warn", "crit", "unknown")]
        bad.sort(key=lambda ch: {"crit": 0, "warn": 1, "unknown": 2}[ch["level"]])
        note = "\n".join(f"[{LEVEL_LABEL[ch['level']]}] {ch['label']}: {ch['summary']}" for ch in bad[:4]) or "모든 항목 정상"
        sh.add([_kind(d), d["name"], d["host"], LEVEL_LABEL[d["status"]],
                _pct(m.get("cpu")), _pct(m.get("mem")), fmt_uptime(m.get("uptime")) if m.get("uptime") is not None else "-",
                d.get("issues", 0), note],
               fmts={5: PCT, 6: PCT, 8: "0"}, level=d["status"], level_col=4, wrap={9})
    sh.finish(f"CPU는 백본 Last 1 min(슬롯 중 최대), 컨트롤러 100% - idle 기준. 판단 기준: CPU 주의 {th['cpu_warn']:g}% / 장애 {th['cpu_crit']:g}%, "
              f"메모리 사용 주의 {th['mem_warn']:g}% / 장애 {th['mem_crit']:g}%.")

    # ------------------------------------------------------------------ 백본 CPU·메모리
    ws = wb.create_sheet("백본 CPU·메모리")
    sh = Sheet(ws, "백본 CPU·메모리 (show cpu summary / show memory summary)", [
        ("장비", 16), ("IP", 15), ("슬롯", 8), ("Last 5 sec", 11), ("Last 1 min", 11), ("Last 5 min", 11),
        ("Memory Total (MB)", 14), ("Used (MB)", 12), ("Free (MB)", 12), ("FreeRatio", 11), ("상태", 10)], sub)
    for d in [x for x in devs if x["role"] == "BB"]:
        m = d.get("metrics", {})
        cpu = m.get("cpu_detail") or {}
        mem = m.get("mem_detail") or {}
        cslots = {str(x["slot"]).split("/")[0]: x for x in cpu.get("slots", [])}
        mslots = {str(x["slot"]).split("/")[0]: x for x in mem.get("slots", [])}
        keys = sorted(set(cslots) | set(mslots), key=lambda k: (len(k), k)) or [""]
        if not d["connected"] or (not cslots and not mslots):
            sh.add([d["name"], d["host"], "-", None, None, None, None, None, None, None, LEVEL_LABEL[d["status"]]],
                   level=d["status"], level_col=11)
            continue
        for k in keys:
            cc, mm = cslots.get(k, {}), mslots.get(k, {})
            free_ratio = mm.get("free_ratio")
            used_pct = None if free_ratio is None else 100 - free_ratio
            lv = _level_for(cc.get("1m"), th["cpu_warn"], th["cpu_crit"])
            lv2 = _level_for(used_pct, th["mem_warn"], th["mem_crit"])
            rank = {"ok": 0, "unknown": 1, "warn": 2, "crit": 3}
            row_lv = max([lv or "unknown", lv2 or "unknown"], key=lambda x: rank[x])
            sh.add([d["name"], d["host"], k or "-", _pct(cc.get("5s")), _pct(cc.get("1m")), _pct(cc.get("5m")),
                    _mb(mm.get("total_kb")), _mb(mm.get("used_kb")), _mb(mm.get("free_kb")), _pct(free_ratio),
                    LEVEL_LABEL[row_lv]],
                   fmts={4: PCT, 5: PCT, 6: PCT, 7: NUM, 8: NUM, 9: NUM, 10: PCT}, level=row_lv, level_col=11)
    sh.finish("상태는 슬롯별 Last 1 min CPU와 메모리 사용률(100% - FreeRatio)을 판단 기준과 비교한 값입니다. "
              "MB = 장비가 보고한 KB ÷ 1024.")

    # ------------------------------------------------------------------ 컨트롤러 CPU·메모리
    ws = wb.create_sheet("컨트롤러 CPU·메모리")
    sh = Sheet(ws, "컨트롤러 CPU·메모리 (show cpu / show memory)", [
        ("구분", 10), ("장비", 16), ("IP", 15), ("모델", 13), ("CPU user", 11), ("CPU system", 11), ("CPU idle", 11),
        ("Memory Total (MB)", 14), ("Used (MB)", 12), ("Free (MB)", 12), ("Used %", 10), ("상태", 10)], sub)
    for d in [x for x in devs if x["role"] in ("MM", "WLC")]:
        m = d.get("metrics", {})
        cpu = m.get("cpu_detail") or {}
        mem = m.get("mem_detail") or {}
        lv = d["status"] if not d["connected"] else max(
            [_level_for(m.get("cpu"), th["cpu_warn"], th["cpu_crit"]) or "unknown",
             _level_for(m.get("mem"), th["mem_warn"], th["mem_crit"]) or "unknown"],
            key=lambda x: {"ok": 0, "unknown": 1, "warn": 2, "crit": 3}[x])
        sh.add([_kind(d), d["name"], d["host"], m.get("model") or "-", _pct(cpu.get("user")), _pct(cpu.get("system")), _pct(cpu.get("idle")),
                _mb(mem.get("total_kb")), _mb(mem.get("used_kb")), _mb(mem.get("free_kb")), _pct(mem.get("used_pct")),
                LEVEL_LABEL[lv]],
               fmts={5: PCT, 6: PCT, 7: PCT, 8: NUM, 9: NUM, 10: NUM, 11: PCT}, level=lv, level_col=12)
    sh.finish("상태는 CPU(100% - idle)와 메모리 Used %를 판단 기준과 비교한 값입니다(장비 전체 상태는 요약 시트 참고).")

    # ------------------------------------------------------------------ 백본 로그
    ws = wb.create_sheet("백본 로그")
    sh = Sheet(ws, "백본 로그 (show logbuffer reverse, 최근순 최대 100건)", [
        ("장비", 16), ("시간", 24), ("등급", 8), ("모듈", 26), ("내용", 90)], sub)
    any_log = False
    for d in [x for x in devs if x["role"] == "BB"]:
        for lg in d.get("metrics", {}).get("log", []):
            any_log = True
            sev = lg.get("severity")
            serious = sev is not None and sev <= 3
            sh.add([d["name"], lg.get("time"), lg.get("severity_name") or "-", lg.get("module"), lg.get("message")],
                   wrap={5}, row_fill="FBE7E5" if serious else None)
    if not any_log:
        sh.add(["로그를 가져오지 못했거나 로그가 없습니다.", None, None, None, None])
    sh.finish("붉은 행: 등급 0~3(긴급·경보·치명·오류). 등급은 Comware 로그 심각도 숫자(0 긴급 ~ 7 디버그) 기준.")

    # ------------------------------------------------------------------ 백본 VRRP
    ws = wb.create_sheet("백본 VRRP")
    sh = Sheet(ws, "백본 VRRP (현재 역할 · 장비 로그의 상태 변경 시각)", [
        ("장비", 16), ("그룹 (인터페이스/VRID)", 22), ("현재 역할", 12), ("점검 결과", 40), ("변경 시각(장비 로그)", 24),
        ("변경", 22), ("장비 로그", 90)], sub)
    any_v = False
    for d in [x for x in devs if x["role"] == "BB"]:
        m = d.get("metrics", {})
        vc = next((c for c in d["checks"] if c["key"] == "vrrp"), None)
        for g, role in (m.get("vrrp") or {}).items():
            any_v = True
            sh.add([d["name"], g, role, (vc or {}).get("summary", ""), None, None, None],
                   level=(vc or {}).get("level"), level_col=4)
        for x in m.get("vrrp_logs") or []:
            any_v = True
            sh.add([d["name"], x.get("group"), None, None, x.get("time"), f"{x.get('old')} → {x.get('new')}", x.get("text")],
                   wrap={7}, row_fill="F6E8CC")
    if not any_v:
        sh.add(["VRRP 정보가 없습니다(장비에 VRRP가 없거나 명령을 가져오지 못함).", None, None, None, None, None, None])
    sh.finish("노란 행: 장비 로그(show logbuffer reverse)에 남은 VRRP 상태 변경. 확인 처리는 프로그램의 ‘VRRP 이력’ 탭에서 합니다.")

    # ------------------------------------------------------------------ 점검 항목
    ws = wb.create_sheet("점검 항목")
    sh = Sheet(ws, "점검 항목별 결과", [("장비", 16), ("구분", 8), ("항목", 16), ("결과", 10), ("내용", 60),
                                        ("근거", 60), ("명령", 30)], sub)
    for d in devs:
        for ch in d["checks"]:
            sh.add([d["name"], _kind(d), ch["label"], LEVEL_LABEL[ch["level"]], ch["summary"],
                    "\n".join(ch.get("evidence", [])), ch.get("command") or ch.get("source", "")],
                   level=ch["level"], level_col=4, wrap={5, 6})
    sh.finish()

    # ------------------------------------------------------------------ 원본 출력
    ws = wb.create_sheet("원본 출력")
    sh = Sheet(ws, "장비 원본 출력 (확인용)", [("장비", 16), ("항목", 16), ("실행 명령", 30), ("결과", 10), ("출력", 110)], sub)
    for d in devs:
        for cmd in d.get("commands", []):
            r = raw(d["id"], cmd["cmd_id"]) or {}
            out = r.get("output", "") or (cmd.get("error") or "")
            sh.add([d["name"], cmd["label"], cmd["command"], "성공" if cmd["ok"] else "실패", out], wrap={5})
            ws.cell(row=sh.row, column=5).font = Font(name="Consolas", size=9)
    sh.finish()

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def report_filename(run: dict) -> str:
    stamp = time.strftime("무선망점검_%Y%m%d_%H%M%S", time.localtime(run["started"]))
    return f"{stamp}_run{int(run['id'])}.xlsx"


def save_report(run: dict, site: str, thresholds: dict, raw: Callable[[str, str], dict | None],
                folder: Path, keep_days: int) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / report_filename(run)
    data = build_workbook(run, site, thresholds, raw)
    with tempfile.NamedTemporaryFile(dir=folder, prefix=path.stem + "_", suffix=".tmp", delete=False) as handle:
        tmp = Path(handle.name)
    try:
        tmp.write_bytes(data)
        tmp.replace(path)
    finally:
        tmp.unlink(missing_ok=True)
    cutoff = time.time() - keep_days * 86400
    for old in folder.glob("무선망점검_*.xlsx"):
        try:
            if old.stat().st_mtime < cutoff:
                old.unlink()
        except OSError:
            pass
    return path
