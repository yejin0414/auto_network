"""Excel headings remove existing illegal controls without changing stored data.

Every assertion reads a saved workbook back from bytes. Only one temporary
synthetic SQLite database is used; no device, server, or external file opens.
"""
from __future__ import annotations

import copy
import io
import tempfile
import time
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from nc.config import DEFAULTS
from nc.excel import FILL, FONT, INK, PCT, Sheet, build_workbook
from nc.report_context import LEGACY_NOTE, capture_context
from nc.store import Store


SHEETS = ["요약", "백본 CPU·메모리", "컨트롤러 CPU·메모리", "백본 로그", "백본 VRRP", "점검 항목", "원본 출력"]
BAD_CONTROLS = "".join(chr(code) for code in range(32) if code not in (9, 10, 13))


def configuration(site):
    cfg = copy.deepcopy(DEFAULTS)
    cfg["site_name"] = site
    return cfg


def synthetic_run(config=None):
    device = {
        "id": "mm", "name": "Synthetic MM", "role": "MM", "host": "192.0.2.10",
        "connected": True, "status": "warn", "issues": 1, "commands": [],
        "metrics": {"cpu": 60, "mem": 25, "uptime": 3600,
                    "cpu_detail": {"user": 60, "system": 0, "idle": 40},
                    "mem_detail": {"total_kb": 1024, "used_kb": 256,
                                   "free_kb": 768, "used_pct": 25}},
        "checks": [
            {"key": "cpu", "label": "CPU", "level": "warn", "summary": "Saved CPU warning", "evidence": []},
            {"key": "memory", "label": "Memory", "level": "ok", "summary": "Saved memory result", "evidence": []},
        ],
    }
    return {"id": 7, "started": 1700000000, "finished": 1700000010, "trigger": "manual",
            "summary": {"overall": "warn", "counts": {"warn": 1}, "total_devices": 1},
            "devices": {"mm": device},
            "report_context": capture_context(config) if config is not None else None}


class ExcelTextSafety(unittest.TestCase):
    def reopen(self, data):
        if isinstance(data, Workbook):
            buffer = io.BytesIO()
            try:
                data.save(buffer)
            finally:
                data.close()
            data = buffer.getvalue()
        workbook = load_workbook(io.BytesIO(data), data_only=False)
        self.addCleanup(workbook.close)
        return workbook

    def report(self, run, cfg):
        return self.reopen(build_workbook(run, cfg["site_name"], cfg["thresholds"], lambda *_: None))

    @staticmethod
    def text(workbook):
        return "\n".join(str(cell.value) for sheet in workbook for row in sheet for cell in row
                         if cell.value is not None)

    def assert_headings(self, workbook, site, run):
        self.assertEqual(workbook.sheetnames, SHEETS)
        self.assertEqual(workbook["요약"]["A1"].value, f"{site} 무선망 점검 보고서")
        stamp = time.strftime("%Y-%m-%d %H:%M", time.localtime(run["started"]))
        subtitle = f"{site}  |  점검 시각 {stamp}  |  회차 #{run['id']}  |  수동 점검"
        for sheet in workbook:
            self.assertEqual(sheet["A2"].value, subtitle)
            self.assertEqual(sheet["A1"].data_type, "s")
            self.assertEqual(sheet["A2"].data_type, "s")

    def test_title_and_subtitle_follow_existing_control_character_policy(self):
        # Tabs, line breaks, Unicode and DEL are outside the existing illegal
        # control range. Compare their serialized form with an ordinary row so
        # XML newline serialization differences do not create platform failures.
        allowed = "한글 café & <text>\tTAB\nLF\rCR\x7fEND"
        dirty = "앞" + BAD_CONTROLS + allowed + "뒤"
        for heading in ("title", "subtitle"):
            with self.subTest(heading=heading):
                workbook = Workbook()
                sheet = Sheet(workbook.active,
                              dirty if heading == "title" else "Normal title",
                              [("Text", 20), ("Allowed text", 20)],
                              dirty if heading == "subtitle" else "Normal subtitle")
                sheet.add([dirty, "앞" + allowed + "뒤"])
                saved = self.reopen(workbook).active
                cell = saved["A1" if heading == "title" else "A2"]
                self.assertEqual(cell.value, saved["A5"].value)
                self.assertEqual(cell.value, saved["B5"].value)
                self.assertEqual(cell.data_type, "s")
                self.assertIn("한글 café & <text>", cell.value)
                self.assertIn("\tTAB\nLF", cell.value)
                self.assertIn("\x7fEND", cell.value)
                self.assertFalse(any(character in cell.value for character in BAD_CONTROLS))

    def test_legacy_site_controls_do_not_break_any_report_sheet_or_mutate_inputs(self):
        cfg = configuration("Current\x00 Campus\x0b & 북부\x1f")
        cfg["thresholds"].update(cpu_warn=80, cpu_crit=90)
        run = synthetic_run()
        before = copy.deepcopy((run, cfg))
        workbook = self.report(run, cfg)
        self.assert_headings(workbook, "Current Campus & 북부", run)
        self.assertEqual(workbook["컨트롤러 CPU·메모리"]["L5"].value, "주의")
        self.assertIn(LEGACY_NOTE, self.text(workbook))
        self.assertEqual((run, cfg), before, "export must not sanitize the original run or current settings")

    def test_saved_site_controls_are_export_only_and_do_not_rewrite_history(self):
        original_cfg = configuration("Saved\x07 Campus\x0c & 남부")
        current_cfg = configuration("Current renamed campus")
        current_cfg["thresholds"].update(cpu_warn=80, cpu_crit=90)
        observation = synthetic_run()
        settings_before = copy.deepcopy((original_cfg, current_cfg))
        with tempfile.TemporaryDirectory() as directory:
            store = Store(Path(directory) / "synthetic.sqlite3")
            try:
                run_id = store.save_run(observation["started"], observation["finished"], "manual",
                                        observation, {}, None, report_config=original_cfg)
                run = store.get_run(run_id)
                run_before = copy.deepcopy(run)
                database_before = tuple(store._db.iterdump())
                changes_before = store._db.total_changes
                workbook = self.report(run, current_cfg)
                self.assert_headings(workbook, "Saved Campus & 남부", run)
                self.assertNotIn(current_cfg["site_name"], self.text(workbook))
                self.assertIn("CPU 주의 50% / 장애 70%", self.text(workbook))
                self.assertEqual(workbook["컨트롤러 CPU·메모리"]["L5"].value, "주의")
                self.assertEqual(run, run_before)
                self.assertEqual(tuple(store._db.iterdump()), database_before)
                self.assertEqual(store._db.total_changes, changes_before)
                self.assertEqual(store.get_run(run_id)["report_context"]["site_name"], original_cfg["site_name"])
            finally:
                store._db.close()
        self.assertEqual((original_cfg, current_cfg), settings_before)

    def test_normal_context_headings_cells_and_formats_are_preserved(self):
        cfg = configuration("Campus 한글 & <A>")
        run = synthetic_run(cfg)
        workbook = self.report(run, cfg)
        self.assert_headings(workbook, cfg["site_name"], run)
        for sheet in workbook:
            self.assertFalse(sheet.sheet_view.showGridLines)
            self.assertEqual(sheet.freeze_panes, "A5")
            self.assertEqual(sheet.row_dimensions[4].height, 30)
            self.assertEqual((sheet["A1"].font.name, sheet["A1"].font.sz, sheet["A1"].font.bold),
                             (FONT, 14, True))
            self.assertEqual(sheet["A1"].font.color.rgb[-6:], INK)
            self.assertEqual((sheet["A2"].font.name, sheet["A2"].font.sz), (FONT, 9))
            self.assertEqual(sheet["A2"].font.color.rgb[-6:], "5D6879")
        summary = workbook["요약"]
        self.assertEqual([cell.value for cell in summary[4]],
                         ["구분", "장비", "IP", "상태", "CPU (%)", "메모리 사용 (%)", "가동시간", "이상 항목 수", "확인할 내용"])
        self.assertEqual(summary.auto_filter.ref, "A4:I5")
        self.assertEqual(summary.column_dimensions["I"].width, 70)
        self.assertEqual((summary["B5"].value, summary["D5"].value), ("Synthetic MM", "주의"))
        self.assertEqual((summary["E5"].value, summary["F5"].value), (0.6, 0.25))
        self.assertEqual((summary["E5"].data_type, summary["F5"].data_type), ("n", "n"))
        self.assertEqual((summary["E5"].number_format, summary["F5"].number_format), (PCT, PCT))
        self.assertEqual(summary["D5"].fill.fgColor.rgb[-6:], FILL["warn"])
        self.assertEqual(workbook["컨트롤러 CPU·메모리"]["L5"].value, "주의")

    def test_formula_like_headings_and_row_values_remain_literal_with_numeric_types(self):
        # Removing a leading illegal control must not turn literal text into a
        # formula; the existing numeric, boolean and blank types also stay intact.
        for prefix in ("", "\x00"):
            with self.subTest(illegal_prefix=bool(prefix)):
                workbook = Workbook()
                sheet = Sheet(workbook.active, prefix + "=SUM(1,2)",
                              [(f"Column {i}", 15) for i in range(1, 9)], prefix + "=2+3")
                values = ["=1+1", "+SUM(1,2)", "-1+2", "@SUM(1,2)", 0, 0.25, True, None]
                sheet.add(values, fmts={6: PCT})
                saved = self.reopen(workbook).active
                for coordinate, text in (("A1", "=SUM(1,2)"), ("A2", "=2+3"),
                                         *zip(("A5", "B5", "C5", "D5"), values[:4])):
                    self.assertEqual(saved[coordinate].value, text)
                    self.assertEqual(saved[coordinate].data_type, "s")
                self.assertEqual((saved["E5"].value, saved["E5"].data_type), (0, "n"))
                self.assertEqual((saved["F5"].value, saved["F5"].data_type, saved["F5"].number_format),
                                 (0.25, "n", PCT))
                self.assertIs(saved["G5"].value, True)
                self.assertEqual(saved["G5"].data_type, "b")
                self.assertIsNone(saved["H5"].value)
                self.assertFalse(any(cell.data_type == "f" for row in saved for cell in row))

    def test_heading_length_preserves_openpyxl_limit_and_existing_row_limit(self):
        # Headings already allow 32,767 characters; body rows have their own
        # 32,000-character limit. The control fix must not shorten legal headings.
        for length in (32001, 32767, 32772):
            for prefix in ("", BAD_CONTROLS):
                with self.subTest(length=length, illegal_prefix=bool(prefix)):
                    clean_text = "한" * length
                    text = prefix + clean_text
                    workbook = Workbook()
                    sheet = Sheet(workbook.active, text, [("Text", 20)], text)
                    sheet.add([text])
                    saved = self.reopen(workbook).active
                    for coordinate in ("A1", "A2"):
                        self.assertEqual(saved[coordinate].value, clean_text[:32767])
                        self.assertEqual(saved[coordinate].data_type, "s")
                    self.assertEqual(saved["A5"].value, clean_text[:32000])
                    self.assertEqual(saved["A5"].data_type, "s")


if __name__ == "__main__":
    unittest.main()
