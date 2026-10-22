"""Historical report metadata, legacy SQLite migration, and all export paths.

Only temporary files/DBs and synthetic results are used. No device or server starts.
"""
from __future__ import annotations

import copy
import io
import json
import sqlite3
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from contextlib import ExitStack, closing
from pathlib import Path
from threading import Barrier
from unittest.mock import Mock, patch

from openpyxl import load_workbook

from nc.config import validate
from nc.credentials import CredentialStore
from nc.engine import Engine
from nc.excel import build_workbook
from nc.report import render
from nc.report_context import EVALUATION_VERSION, LEGACY_NOTE, capture_context
from nc.server import App, make_handler
from nc.store import Store


def configuration(site='Synthetic original site', warn=50, crit=70):
    return validate({'site_name': site, 'thresholds': {'cpu_warn': warn, 'cpu_crit': crit},
                     'report': {'excel_auto': True, 'notify': False}})


def evaluation():
    checks = [
        {'key': 'cpu', 'label': 'CPU', 'level': 'warn', 'summary': 'Saved CPU warning', 'evidence': []},
        {'key': 'memory', 'label': 'Memory', 'level': 'ok', 'summary': 'Saved memory normal', 'evidence': []},
    ]
    devices = {}
    for role in ['BB', 'WLC']:
        devices[role] = {
            'id': role, 'name': role, 'host': '192.0.2.1' if role == 'BB' else '192.0.2.2',
            'role': role, 'connected': True, 'status': 'warn', 'issues': 1,
            'checks': copy.deepcopy(checks),
            'metrics': {'cpu': 60, 'mem': 20,
                        'cpu_detail': {'user': 60, 'idle': 40, 'system': 0,
                                       'slots': [{'slot': '1', '1m': 60, '5s': 60, '5m': 60}]},
                        'mem_detail': {'used_pct': 20, 'total_kb': 1000, 'used_kb': 200, 'free_kb': 800,
                                       'slots': [{'slot': '1', 'free_ratio': 80,
                                                  'total_kb': 1000, 'used_kb': 200, 'free_kb': 800}]}}
        }
    return {'summary': {'overall': 'warn', 'counts': {'warn': 2}, 'total_devices': 2}, 'devices': devices}


def workbook_text(wb):
    return '\n'.join(str(cell.value) for ws in wb for row in ws for cell in row if cell.value is not None)


class SnapshotDatabase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / 'history.sqlite3'
        self.store = Store(self.path)
        self.addCleanup(self.store._db.close)

    def test_only_allowlisted_metadata_survives_reopen_and_mutation(self):
        cfg = configuration()
        cfg['password'] = 'must-not-persist'
        cfg['credentials'] = {'password': 'must-not-persist'}
        cfg['thresholds']['api_token'] = 'must-not-persist'
        cfg['devices'] = [{'host': 'private-operational-value'}]
        rid = self.store.save_run(1700000000, 1700000010, 'manual', evaluation(), {}, None, report_config=cfg)
        cfg['site_name'] = 'Changed later'
        cfg['thresholds']['cpu_warn'] = 90
        stored = self.store.get_run(rid)
        self.assertEqual(stored['report_context']['site_name'], 'Synthetic original site')
        self.assertEqual(stored['report_context']['thresholds']['cpu_warn'], 50)
        self.assertEqual(stored['report_context']['evaluation_version'], EVALUATION_VERSION)
        self.assertEqual(set(stored['report_context']), {'schema_version', 'evaluation_version', 'site_name', 'thresholds'})
        self.assertEqual(len(stored['report_context']['thresholds']), 6)
        context_json = self.store._db.execute('SELECT report_context FROM runs WHERE id=?', (rid,)).fetchone()[0]
        self.assertNotIn('must-not-persist', context_json)
        self.assertNotIn('private-operational-value', context_json)
        stored['report_context']['thresholds']['cpu_warn'] = 1
        self.store._db.close()
        reopened = Store(self.path)
        self.addCleanup(reopened._db.close)
        self.assertEqual(reopened.get_run(rid)['report_context']['thresholds']['cpu_warn'], 50)
        self.assertEqual(reopened.get_run(rid)['devices']['WLC']['checks'][0]['level'], 'warn')

    def test_old_save_run_call_remains_supported(self):
        rid = self.store.save_run(1700000000, 1700000010, 'manual', evaluation(), {}, None)
        self.assertIsNone(self.store.get_run(rid)['report_context'])
        self.assertEqual(self.store.list_runs()[0]['id'], rid)

    def test_broken_or_future_snapshot_never_becomes_current_thresholds(self):
        rid = self.store.save_run(1700000000, 1700000010, 'manual', evaluation(), {}, None)
        cases = ['not-json', '{}', json.dumps({'schema_version': 999}),
                 json.dumps({**capture_context(configuration()), 'evaluation_version': '999'}),
                 json.dumps({**capture_context(configuration()), 'thresholds': {'cpu_warn': 'nan'}})]
        for raw in cases:
            with self.subTest(raw=raw):
                with self.store._db:
                    self.store._db.execute('UPDATE runs SET report_context=? WHERE id=?', (raw, rid))
                run = self.store.get_run(rid)
                self.assertIsNone(run['report_context'])
                wb = load_workbook(io.BytesIO(build_workbook(run, 'Current', configuration(warn=80, crit=90)['thresholds'], lambda *_: None)))
                self.assertEqual(wb['컨트롤러 CPU·메모리']['L5'].value, '주의')
                self.assertIn('당시 판단 기준 없음', workbook_text(wb))


class LegacyDatabaseMigration(unittest.TestCase):
    def legacy_file(self, path):
        with closing(sqlite3.connect(path)) as db, db:
            db.execute('CREATE TABLE runs (id INTEGER PRIMARY KEY AUTOINCREMENT, started REAL NOT NULL, '
                       'finished REAL, trigger TEXT, overall TEXT, summary TEXT)')
            db.execute('CREATE TABLE device_results (run_id INTEGER NOT NULL, device_id TEXT NOT NULL, '
                       'status TEXT, cpu REAL, mem REAL, data TEXT, PRIMARY KEY (run_id, device_id))')
            data = evaluation()
            db.execute('INSERT INTO runs VALUES (1,1700000000,1700000010,?,?,?)',
                       ('manual', 'warn', json.dumps(data['summary'])))
            for key, dev in data['devices'].items():
                db.execute('INSERT INTO device_results VALUES (?,?,?,?,?,?)',
                           (1, key, dev['status'], 60, 20, json.dumps(dev)))

    def test_existing_history_migrates_once_and_remains_readable(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'legacy.sqlite3'
            self.legacy_file(path)
            for _ in range(3):
                store = Store(path)
                try:
                    old = store.get_run(1)
                    self.assertEqual(old['summary']['overall'], 'warn')
                    self.assertEqual(len(old['devices']), 2)
                    self.assertIsNone(old['report_context'])
                    columns = [r['name'] for r in store._db.execute('PRAGMA table_info(runs)')]
                    self.assertEqual(columns.count('report_context'), 1)
                finally:
                    store._db.close()
            store = Store(path)
            try:
                new = store.save_run(1700001000, 1700001010, 'manual', evaluation(), {}, None,
                                     report_config=configuration())
                self.assertEqual(new, 2)
                self.assertIsNone(store.get_run(1)['report_context'])
                self.assertEqual(store.get_run(2)['report_context']['thresholds']['cpu_warn'], 50)
            finally:
                store._db.close()

    def test_startup_recovers_after_uncommitted_schema_upgrade(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'interrupted.sqlite3'
            self.legacy_file(path)
            db = sqlite3.connect(path)
            db.execute('BEGIN IMMEDIATE')
            db.execute('ALTER TABLE runs ADD COLUMN report_context TEXT')
            db.close()  # Simulates exit before the migration transaction commits.
            store = Store(path)
            try:
                self.assertEqual(store.get_run(1)['overall'], 'warn')
                self.assertIsNone(store.get_run(1)['report_context'])
            finally:
                store._db.close()

    def test_independent_store_openers_do_not_duplicate_schema_upgrade(self):
        with tempfile.TemporaryDirectory() as tmp:
            for legacy in [False, True]:
                with self.subTest(legacy=legacy):
                    path = Path(tmp) / f'concurrent-{legacy}.sqlite3'
                    if legacy:
                        self.legacy_file(path)
                    ready = Barrier(3, timeout=5)

                    def open_read_close(_):
                        ready.wait()
                        store = Store(path)
                        try:
                            columns = [r['name'] for r in store._db.execute('PRAGMA table_info(runs)')]
                            self.assertEqual(columns.count('report_context'), 1)
                            run = store.get_run(1)
                            return run['summary']['overall'] if run else None
                        finally:
                            store._db.close()

                    with ThreadPoolExecutor(max_workers=3) as pool:
                        self.assertEqual(list(pool.map(open_read_close, range(3))),
                                         ['warn' if legacy else None] * 3)


class _InitializationConnection(sqlite3.Connection):
    """Inject SQLite errors while retaining a real temporary database handle."""
    before_execute = None
    closed = False

    def execute(self, sql, parameters=()):
        if self.before_execute:
            self.before_execute(sql)
        return super().execute(sql, parameters)

    def close(self):
        super().close()
        self.closed = True


class StoreInitializationRetries(unittest.TestCase):
    @staticmethod
    def database_error(code, message='database is locked'):
        error = sqlite3.OperationalError(message)
        error.sqlite_errorcode = code
        return error

    def test_transient_busy_during_journal_selection_or_schema_retries_successfully(self):
        for statement, code in [('PRAGMA journal_mode=', sqlite3.SQLITE_BUSY),
                                ('ALTER TABLE', sqlite3.SQLITE_LOCKED)]:
            with self.subTest(statement=statement), tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / 'transient.sqlite3'
                LegacyDatabaseMigration().legacy_file(path)
                with closing(sqlite3.connect(path, timeout=0, factory=_InitializationConnection)) as db:
                    failures = []

                    def fail_once(sql):
                        if sql.startswith(statement) and not failures:
                            failures.append(sql)
                            raise self.database_error(code)

                    db.before_execute = fail_once
                    with patch('nc.store.sqlite3.connect', return_value=db), patch('nc.store.time.sleep') as sleep:
                        store = Store(path)
                    self.assertEqual(len(failures), 1)
                    sleep.assert_called_once()
                    self.assertEqual(store.get_run(1)['overall'], 'warn')
                    columns = [r['name'] for r in db.execute('PRAGMA table_info(runs)')]
                    self.assertEqual(columns.count('report_context'), 1)
                    self.assertEqual(db.execute('PRAGMA busy_timeout').fetchone()[0], 30000)

    def test_persistent_busy_stops_at_deadline_and_closes_connection(self):
        with tempfile.TemporaryDirectory() as tmp, ExitStack() as cleanup:
            path = Path(tmp) / 'busy.sqlite3'
            db = sqlite3.connect(path, timeout=0, factory=_InitializationConnection)
            cleanup.callback(db.close)
            elapsed = [0.0]
            busy = self.database_error(sqlite3.SQLITE_BUSY | 0x100)
            db.before_execute = Mock(side_effect=busy)

            def advance(seconds):
                elapsed[0] += seconds

            with patch('nc.store.sqlite3.connect', return_value=db), \
                    patch('nc.store._INIT_TIMEOUT', 0.1), \
                    patch('nc.store.time.monotonic', side_effect=lambda: elapsed[0]), \
                    patch('nc.store.time.sleep', side_effect=advance) as sleep:
                with self.assertRaises(sqlite3.OperationalError) as raised:
                    Store(path)
            self.assertIs(raised.exception, busy)
            self.assertEqual(elapsed[0], 0.1)
            self.assertEqual(sleep.call_count, 2)
            self.assertEqual(db.before_execute.call_count, 3)
            self.assertTrue(db.closed)
            with self.assertRaises(sqlite3.ProgrammingError):
                db.cursor()
            # Keep the exception/connection referenced while unlinking: Windows
            # must not need garbage collection to release the failed startup.
            path.unlink()

    def test_nonbusy_failure_rolls_back_schema_and_closes_without_retry(self):
        with tempfile.TemporaryDirectory() as tmp, ExitStack() as cleanup:
            path = Path(tmp) / 'invalid.sqlite3'
            LegacyDatabaseMigration().legacy_file(path)
            db = sqlite3.connect(path, timeout=0, factory=_InitializationConnection)
            cleanup.callback(db.close)
            error = self.database_error(sqlite3.SQLITE_ERROR, 'injected schema failure')

            def fail_alter(sql):
                if sql.startswith('ALTER TABLE'):
                    self.assertTrue(db.in_transaction)
                    raise error

            db.before_execute = fail_alter
            with patch('nc.store.sqlite3.connect', return_value=db), patch('nc.store.time.sleep') as sleep:
                with self.assertRaises(sqlite3.OperationalError) as raised:
                    Store(path)
            self.assertIs(raised.exception, error)
            sleep.assert_not_called()
            self.assertTrue(db.closed)
            with self.assertRaises(sqlite3.ProgrammingError):
                db.cursor()
            with closing(sqlite3.connect(path)) as reopened:
                tables = {row[0] for row in reopened.execute("SELECT name FROM sqlite_master WHERE type='table'")}
                self.assertEqual(tables, {'runs', 'device_results', 'sqlite_sequence'})
                self.assertEqual(reopened.execute('SELECT COUNT(*) FROM runs').fetchone()[0], 1)
                self.assertNotIn('report_context', [r[1] for r in reopened.execute('PRAGMA table_info(runs)')])


class ReportSnapshotExports(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.folder = Path(self.tmp.name)
        self.store = Store(self.folder / 'db.sqlite3')
        self.addCleanup(self.store._db.close)
        self.old_cfg = configuration()
        self.current_cfg = configuration('Current renamed site', 80, 90)
        self.rid = self.store.save_run(1700000000, 1700000010, 'manual', evaluation(), {}, None,
                                      report_config=self.old_cfg)
        self.legacy_rid = self.store.save_run(1700000020, 1700000030, 'manual', evaluation(), {}, None)
        cfg = Mock()
        cfg.get.return_value = self.current_cfg
        self.engine = Engine(cfg, self.store, CredentialStore(persistent=False), Mock(), demo=True)
        self.engine.reports_dir = self.folder / 'reports'
        self.app = App(self.engine, self.folder, version='test')

    def assert_original_workbook(self, data):
        wb = load_workbook(io.BytesIO(data))
        self.assertIn('Synthetic original site', wb['요약']['A1'].value)
        self.assertNotIn('Current renamed site', workbook_text(wb))
        self.assertEqual(wb['백본 CPU·메모리']['K5'].value, '주의')
        self.assertEqual(wb['컨트롤러 CPU·메모리']['L5'].value, '주의')
        self.assertIn('CPU 주의 50% / 장애 70%', workbook_text(wb))
        return wb

    def get_response(self, path):
        handler = object.__new__(make_handler(self.app))
        handler.path = path
        handler.headers = {'Host': 'localhost:8875'}
        handler._send = Mock()
        handler.do_GET()
        self.assertEqual(handler._send.call_args.args[0], 200)
        return handler._send.call_args.args[1]

    def test_excel_and_html_direct_render_use_the_stored_context(self):
        run = self.store.get_run(self.rid)
        self.assert_original_workbook(build_workbook(run, self.current_cfg['site_name'], self.current_cfg['thresholds'], lambda *_: None))
        html = render(run, self.current_cfg['site_name'])
        self.assertIn('Synthetic original site', html)
        self.assertNotIn('Current renamed site', html)
        self.assertIn('CPU 주의 50% / 장애 70%', html)
        self.assertIn('Saved CPU warning', html)

    def test_excel_and_html_download_routes_use_the_stored_context(self):
        suffix = f'?id={self.rid}&t={self.app.token}'
        self.assert_original_workbook(self.get_response('/report.xlsx' + suffix))
        for download in ['', '&download=1']:
            html = self.get_response('/report' + suffix + download)
            self.assertIn('Synthetic original site', html)
            self.assertNotIn('Current renamed site', html)
            self.assertIn('CPU 주의 50% / 장애 70%', html)

    def test_save_open_and_automatic_report_paths_use_the_stored_context(self):
        for route in ['/api/excel_save', '/api/excel_open']:
            with self.subTest(route=route), patch('nc.server._open_path') as open_path:
                status, payload = self.app.handle_post(route, {'id': self.rid})
                self.assertEqual(status, 200)
                self.assert_original_workbook(Path(payload['path']).read_bytes())
                self.assertEqual(open_path.call_count, int(route.endswith('open')))
        with patch('nc.notify.show_toast') as notify:
            self.engine._after_run(self.current_cfg, self.rid)
            notify.assert_not_called()
        self.assert_original_workbook(Path(self.engine.last_report).read_bytes())

    def test_legacy_reports_preserve_saved_results_and_disclose_missing_context(self):
        run = self.store.get_run(self.legacy_rid)
        wb = load_workbook(io.BytesIO(build_workbook(run, self.current_cfg['site_name'], self.current_cfg['thresholds'], lambda *_: None)))
        self.assertEqual(wb['백본 CPU·메모리']['K5'].value, '주의')
        self.assertEqual(wb['컨트롤러 CPU·메모리']['L5'].value, '주의')
        self.assertIn(LEGACY_NOTE, workbook_text(wb))
        self.assertIn('슬롯별 재판정은 하지 않았습니다', workbook_text(wb))
        self.assertNotIn('CPU 주의 80%', workbook_text(wb))
        html = render(run, self.current_cfg['site_name'])
        self.assertIn(LEGACY_NOTE, html)
        self.assertIn('Saved CPU warning', html)
        self.assertNotIn('CPU 주의 80%', html)

    def test_legacy_missing_saved_metric_result_stays_unknown(self):
        run = self.store.get_run(self.legacy_rid)
        run['devices']['BB']['checks'] = []
        run['devices']['WLC']['checks'] = []
        wb = load_workbook(io.BytesIO(build_workbook(run, 'Current', self.current_cfg['thresholds'], lambda *_: None)))
        self.assertEqual(wb['백본 CPU·메모리']['K5'].value, '확인 불가')
        self.assertEqual(wb['컨트롤러 CPU·메모리']['L5'].value, '확인 불가')


if __name__ == '__main__':
    unittest.main()
