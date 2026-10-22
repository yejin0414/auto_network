"""Reports retain per-command origins when cluster data comes from two WLCs."""
import io
import unittest

from openpyxl import load_workbook

from nc.config import DEFAULTS
from nc.excel import build_workbook
from nc.report import render


class ClusterSourceReports(unittest.TestCase):
    def run_data(self):
        checks = [{"key": key, "label": key, "summary": "Synthetic result", "level": "ok",
                   "evidence": [], "command": command, "cmd_id": key, "source": name, "source_id": source_id}
                  for key, command, name, source_id in [
                      ("cluster", "show lc-cluster group-membership", "WLC <A>", "a"),
                      ("clients", "show lc-cluster load distribution client", "WLC <B>", "b"),
                  ]]
        return {"id": 1, "started": 1700000000, "trigger": "manual",
                "summary": {"overall": "ok", "counts": {"ok": 1}},
                "devices": {"a": {"id": "a", "name": "WLC A", "role": "WLC", "host": "192.0.2.1",
                                  "status": "ok", "connected": True, "metrics": {}, "checks": checks}}}

    def test_html_shows_escaped_source_beside_each_command(self):
        html = render(self.run_data(), "Synthetic site")
        for command, source in [("group-membership", "A"), ("load distribution client", "B")]:
            self.assertIn(f"show lc-cluster {command}", html)
            self.assertIn(f"출처: WLC &lt;{source}&gt;", html)
        self.assertNotIn("WLC <A>", html)

    def test_excel_keeps_each_command_and_actual_source_together(self):
        data = build_workbook(self.run_data(), "Synthetic site", DEFAULTS["thresholds"], lambda *_: None)
        workbook = load_workbook(io.BytesIO(data))
        try:
            values = [cell.value for row in workbook["점검 항목"] for cell in row]
            self.assertIn("show lc-cluster group-membership\n출처: WLC <A>", values)
            self.assertIn("show lc-cluster load distribution client\n출처: WLC <B>", values)
            self.assertEqual(len(workbook.sheetnames), 7)
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
