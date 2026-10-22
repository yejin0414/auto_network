"""VRRP 변경 이력(언제 바뀌었는지) + 확인·초기화, 응답 없는 장비가 있어도 점검이 끝나는지."""
import json
import tempfile
import threading
import time
import unittest
from pathlib import Path

from nc import parsers as P
from nc.config import ConfigStore
from nc.credentials import Credential, CredentialStore
from nc.demo import DEMO_CONFIG, VRRP_SWAP_AT, DemoWorld, seed_history
from nc.engine import Engine
from nc.excel import build_workbook
from nc.server import App
from nc.store import Store


def make(tmp: Path, mutate=None):
    c = json.loads(json.dumps(DEMO_CONFIG))
    c.setdefault("report", {})["notify"] = False
    if mutate:
        mutate(c)
    (tmp / "config.json").write_text(json.dumps(c, ensure_ascii=False), encoding="utf-8")
    cfg = ConfigStore(tmp / "config.json")
    creds = CredentialStore(persistent=False)
    creds.set("shared", Credential("u", "p"), remember=False)
    w = DemoWorld(speed=0.0)
    w.devices_provider = lambda: cfg.get()["devices"]
    eng = Engine(cfg, Store(tmp / "db.sqlite3"), creds, w, demo=True)
    eng.reports_dir = tmp / "reports"
    eng.demo_quiet = True
    return eng, w


def vrrp_check(run, dev):
    return next(c for c in run["devices"][dev]["checks"] if c["key"] == "vrrp")


class VrrpLogParse(unittest.TestCase):
    def test_comware7_and_5(self):
        t = ("%Oct 22 10:15:32:123 2026 BB-03 VRRP4/6/VRRP_STATUS_CHANGE: The status of IPv4 virtual router 20 "
             "(configured on Vlan-interface20) changed from Master to Backup: Received a VRRP packet with higher priority.\n"
             "%Oct 22 09:00:00:000 2026 BB-03 SHELL/5/SHELL_LOGIN: op logged in.\n"
             "%Apr 26 12:00:00:123 2025 H3C VRRP/6/VRRP_STATE_CHANGE: Interface Vlan-interface10, VRID 1, state from Backup to Master.")
        v = P.parse_vrrp_logs(t)
        self.assertEqual([(x["group"], x["old"], x["new"]) for x in v], [("Vlan20/20", "Master", "Backup"), ("Vlan10/1", "Backup", "Master")])
        self.assertIsNotNone(v[0]["ts"])
        self.assertEqual(time.localtime(v[0]["ts"]).tm_hour, 10)


class VrrpHistory(unittest.TestCase):
    def test_history_ack_and_clear(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as t:
            eng, w = make(Path(t))
            seed_history(eng, w, runs=VRRP_SWAP_AT + 2)
            st = eng.store
            items = st.list_vrrp()
            got = {(x["device_id"], x["grp"], x["old"], x["new"]) for x in items}
            self.assertIn(("bb3", "Vlan20/20", "Master", "Backup"), got)
            self.assertIn(("bb4", "Vlan20/20", "Backup", "Master"), got)
            bb3 = next(x for x in items if x["device_id"] == "bb3")
            self.assertEqual(bb3["source"], "check+log")          # 점검 + 장비 로그로 정확한 시각
            self.assertAlmostEqual(bb3["log_ts"], w.vrrp_ts, delta=2)
            self.assertEqual(len([x for x in items if x["device_id"] == "bb3"]), 1)  # 다음 점검에서 중복으로 쌓이지 않음
            self.assertEqual(st.vrrp_unacked(), 2)
            run = st.get_run()
            self.assertEqual(vrrp_check(run, "bb3")["level"], "warn")      # 기준값 대비 역할 변화 표시 중
            self.assertEqual(vrrp_check(run, "bb4")["level"], "warn")
            bb1_base = st.baselines()["bb1"]
            # 확인 · 초기화 (화면 버튼과 같은 API)
            app = App(eng, Path(t))
            code, r = app.handle_post("/api/vrrp/ack", {"device": None, "recheck": False})
            self.assertEqual(code, 200)
            self.assertIn("bb3", r["devices"])
            self.assertEqual(app.handle_get("/api/vrrp", {})[1]["unacked"], 0)
            self.assertEqual(st.vrrp_unacked(), 0)
            self.assertEqual(st.baselines()["bb1"].get("ospf"), bb1_base.get("ospf"))  # 다른 기준값은 그대로
            self.assertNotIn("vrrp", st.baselines()["bb3"])
            rid = eng.run_blocking("manual")
            run = st.get_run(rid)
            self.assertEqual(vrrp_check(run, "bb3")["level"], "ok")        # 다시 점검하면 VRRP 장애 표시 사라짐
            self.assertEqual(vrrp_check(run, "bb4")["level"], "ok")
            self.assertEqual(st.vrrp_unacked(), 0)                          # 같은 로그가 다시 미확인으로 생기지 않음
            self.assertEqual(st.baselines()["bb3"]["vrrp"]["Vlan20/20"], "Backup")
            import io
            from openpyxl import load_workbook
            data = build_workbook(run, "", eng.cfg.get()["thresholds"], lambda dev, cmd: st.get_raw(rid, dev, cmd))
            wb = load_workbook(io.BytesIO(data))
            self.assertIn("백본 VRRP", wb.sheetnames)
            vals = [c.value for row in wb["백본 VRRP"].iter_rows() for c in row]
            self.assertIn("Master → Backup", vals)


class HungDevice(unittest.TestCase):
    def test_run_finishes_when_device_hangs(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as t:
            eng, w = make(Path(t), lambda c: c.setdefault("ssh", {}).update(device_limit_minutes=0.03))
            eng.device_grace = 1.0
            gate = threading.Event()
            real_open = w.open

            class Hang:
                def __init__(self, s): self.s = s
                def run(self, cmd, timeout): gate.wait(60); return self.s.run(cmd, timeout)
                def close(self): pass
                def abort(self): gate.set()

            w.open = lambda dev, cred, ssh: Hang(real_open(dev, cred, ssh)) if dev["id"] == "bb2" else real_open(dev, cred, ssh)
            t0 = time.time()
            rid = eng.run_blocking("auto")
            self.assertLess(time.time() - t0, 20)
            self.assertFalse(eng.running)
            run = eng.store.get_run(rid)
            res = run["devices"]["bb2"]
            self.assertTrue(gate.is_set())   # 멈춘 세션을 강제로 끊음
            joined = json.dumps(res, ensure_ascii=False)
            self.assertTrue("시간 초과" in joined or "응답 없음" in joined, joined[:500])
            self.assertTrue(run["devices"]["bb1"]["connected"])


if __name__ == "__main__":
    unittest.main()
