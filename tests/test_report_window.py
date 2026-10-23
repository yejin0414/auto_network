"""새 명령 파서, 엑셀 보고서, 완료 알림 문구, 프로그램 창 종료 판단."""
import io
import json
import subprocess
import sys
import tempfile
import time
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from nc import parsers as P
from nc.appwindow import WindowWatch
from nc.config import ConfigStore
from nc.credentials import Credential, CredentialStore
from nc.demo import DEMO_CONFIG, DemoWorld
from nc.engine import Engine
from nc.excel import build_workbook, save_report
from nc.notify import build_script, completion_message
from nc.store import Store

FX = Path(__file__).parent / "fixtures"
r = lambda f: (FX / f).read_text(encoding="utf-8")


class NewCommandParsers(unittest.TestCase):
    def test_cpu_summary(self):
        p = P.parse_cpu_comware(r("comware_cpu_summary.txt"))
        self.assertEqual([(s["5s"], s["1m"], s["5m"]) for s in p["slots"]], [(6, 7, 6), (12, 11, 9)])
        self.assertEqual((p["5s"], p["1m"], p["5m"]), (12, 11, 9))

    def test_memory_summary_free(self):
        p = P.parse_mem_comware(r("comware_mem_summary.txt"))
        self.assertEqual([(s["slot"], s["free_kb"], s["free_ratio"]) for s in p["slots"]],
                         [("1", 6514828, 81.2), ("2", 2514828, 31.4)])
        self.assertEqual(p["used_pct"], 68.6)

    def test_controller_user_system_used(self):
        c = P.parse_cpu_aruba(r("aos8_show_cpu.txt"))
        self.assertEqual((c["user"], c["system"], c["idle"]), (7.2, 9.1, 83.7))
        m = P.parse_mem_aruba(r("aos8_memory.txt"))
        self.assertEqual((m["used_kb"], m["used_pct"]), (1872100, 47.4))

    def test_logbuffer(self):
        logs = P.parse_logbuffer(r("comware_logbuffer_reverse.txt"))
        self.assertEqual(len(logs), 4)
        self.assertEqual(logs[1]["severity"], 3)
        self.assertIn("Ten-GigabitEthernet1/0/2", logs[1]["message"])


class ExcelReport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        temporary = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        cls.addClassCleanup(temporary.cleanup)
        d = Path(temporary.name)
        config = json.loads(json.dumps(DEMO_CONFIG))
        config.setdefault("report", {})["notify"] = False
        (d / "c.json").write_text(json.dumps(config))
        cr = CredentialStore(persistent=False)
        for p in ("backbone", "aruba"):
            cr.set(p, Credential("d", "d"), False)
        cls.engine = Engine(ConfigStore(d / "c.json"), Store(d / "db"), cr, DemoWorld(speed=0), demo=True)
        cls.addClassCleanup(cls.engine.store._db.close)
        cls.engine.reports_dir = d / "보고서"
        # 낮의 균형 분배를 검증하며, 실행 당일을 사용해 이력 보관 기간도 유지한다.
        first_run = datetime.now().replace(hour=13, minute=0, second=0, microsecond=0).timestamp()
        for i in range(27):  # 27번째 회차: 백본 3호기 장애 시나리오
            cls.engine.run_blocking("auto", now=first_run + i * 60)
        cls.run_ = cls.engine.store.get_run()
        cls.dir = d

    def wb(self):
        e, run = self.engine, self.run_
        data = build_workbook(run, "데모", e.cfg.get()["thresholds"], lambda dev, cmd: e.store.get_raw(run["id"], dev, cmd))
        return load_workbook(io.BytesIO(data))

    def test_sheets_and_values(self):
        wb = self.wb()
        self.assertEqual(wb.sheetnames, ["요약", "백본 CPU·메모리", "컨트롤러 CPU·메모리", "백본 로그", "백본 VRRP", "점검 항목", "원본 출력"])
        bb = wb["백본 CPU·메모리"]
        head = [c.value for c in bb[4]]
        self.assertEqual(head[:10], ["장비", "IP", "슬롯", "Last 5 sec", "Last 1 min", "Last 5 min",
                                     "Memory Total (MB)", "Used (MB)", "Free (MB)", "FreeRatio"])
        rows = [list(r) for r in bb.iter_rows(min_row=5, max_row=12, values_only=True)]
        self.assertEqual([x[0] for x in rows], [f"백본 {i}호기" for i in (1, 1, 2, 2, 3, 3, 4, 4)])  # 백본 4대 × 슬롯 2
        cpu = self.run_["devices"]["bb3"]["metrics"]["cpu_detail"]["slots"][0]
        bb3 = next(x for x in rows if x[0] == "백본 3호기")
        self.assertAlmostEqual(bb3[4], cpu["1m"] / 100)
        self.assertEqual(bb.cell(row=5, column=5).number_format, "0.0%")
        ct = wb["컨트롤러 CPU·메모리"]
        self.assertEqual([c.value for c in ct[4]][3:7], ["모델", "CPU user", "CPU system", "CPU idle"])
        w = self.run_["devices"]["wlc1"]["metrics"]
        row = next(r for r in ct.iter_rows(min_row=5, values_only=True) if r[1] == "WLC-IN-01")
        self.assertEqual((row[0], row[3]), ("WLC IN", "Aruba7240XM"))
        self.assertAlmostEqual(row[4], w["cpu_detail"]["user"] / 100)
        self.assertAlmostEqual(row[8], round(w["mem_detail"]["used_kb"] / 1024, 1))
        out = next(r for r in ct.iter_rows(min_row=5, values_only=True) if r[1] == "WLC-OUT-01")
        self.assertEqual((out[0], out[3]), ("WLC OUT", "Aruba7210"))
        logs = list(wb["백본 로그"].iter_rows(min_row=5, values_only=True))
        self.assertTrue(any(x[2] == "오류" and "Ten-GigabitEthernet1/0/2" in (x[4] or "") for x in logs))
        summary = {r[1]: r[3] for r in wb["요약"].iter_rows(min_row=5, values_only=True) if r[1]}
        self.assertEqual(summary["백본 3호기"], "장애")
        self.assertEqual(len([n for n in summary if n.startswith("WLC-")]), 8)

    def test_two_clusters_separate(self):
        s = self.run_["summary"]
        self.assertEqual([(c["name"], c["members"], c["model"]) for c in s["clusters"]],
                         [("IN", 4, "Aruba7240XM"), ("OUT", 4, "Aruba7210")])
        self.assertEqual(s["total_active"], sum(c["total_active"] for c in s["clusters"]))
        out = [d for d in self.run_["devices"].values() if d.get("cluster") == "OUT"]
        # OUT은 OUT끼리만 비교: 수용량이 작아도 'Client 낮음' 주의가 나오지 않음
        self.assertFalse(any(c["key"] == "clients" and c["level"] == "warn" for d in out for c in d["checks"]))
        self.assertTrue(all(any(c["key"] == "cluster" and c["level"] == "ok" for c in d["checks"]) for d in out))
        self.assertTrue(any(c["key"] == "clients" and c["level"] == "warn"
                            for c in self.run_["devices"]["wlc3"]["checks"]))

    def test_auto_saved_after_each_check(self):
        files = list((self.dir / "보고서").glob("무선망점검_*.xlsx"))
        self.assertTrue(files, "점검 후 엑셀이 자동 저장되지 않았습니다")
        load_workbook(files[0])

    def test_completion_message(self):
        title, lines = completion_message(self.run_)
        self.assertEqual(title, "점검 완료: 장애 발생")
        self.assertIn("장애 1", lines[0])
        self.assertIn("백본 3호기", lines[1])
        script = build_script(title, lines + ["it's <ok> & done"])
        self.assertIn("ToastNotificationManager", script)
        self.assertNotIn("<ok>", script)


class WindowWatchLogic(unittest.TestCase):
    def test_bye_then_exit(self):
        closed = []
        w = WindowWatch(None, lambda: closed.append(1))
        w.heartbeat()
        w.last_seen -= 10  # 마지막 신호 뒤에 창을 닫음
        w.bye()
        w.bye_at -= 5
        self.assertTrue(w._closed())

    def test_refresh_is_not_close(self):
        w = WindowWatch(None, lambda: None)
        w.heartbeat()
        w.bye()
        w.bye_at -= 5
        w.heartbeat()  # 새로고침: 바로 다시 신호
        self.assertFalse(w._closed())

    def test_browser_process_exit(self):
        p = subprocess.Popen([sys.executable, "-c", "pass"])
        p.wait()
        w = WindowWatch(p, lambda: None)
        w.heartbeat()
        self.assertFalse(w._closed())
        w.last_seen -= 31
        self.assertTrue(w._closed())

    def test_minimized_window_keeps_running(self):
        p = subprocess.Popen([sys.executable, "-c", "import time; time.sleep(5)"])
        try:
            w = WindowWatch(p, lambda: None)
            w.heartbeat()
            w.last_seen -= 600  # 최소화로 신호가 늦어도 창(프로세스)이 살아 있으면 유지
            self.assertFalse(w._closed())
        finally:
            p.kill()
            p.wait(10)


if __name__ == "__main__":
    unittest.main()
