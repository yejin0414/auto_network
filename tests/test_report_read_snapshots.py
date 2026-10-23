"""Report entry points use one captured run and its raw output during pruning.

All storage and reports are temporary. Handlers are called without listeners;
notifications and opening files are mocked, and no device collection runs.
"""
from __future__ import annotations

import copy
import io
import sqlite3
import tempfile
import time
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import load_workbook

from nc.collector import CmdResult, DeviceCollection
from nc.config import validate
from nc.credentials import CredentialStore
from nc.engine import Engine
from nc.excel import build_workbook
from nc.notify import completion_message
from nc.server import App, make_handler
from nc.store import Store


class ReportReadSnapshots(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.store = Store(self.root / "synthetic.sqlite3")
        self.addCleanup(self.store._db.close)
        self.original_cfg = validate({
            "site_name": "Historical synthetic site",
            "report": {"excel_auto": True, "notify": False},
        })
        self.current_cfg = copy.deepcopy(self.original_cfg)
        self.current_cfg["site_name"] = "Current changed site"
        self.current_cfg["thresholds"].update(cpu_warn=80, cpu_crit=90)
        self.engine = Engine(SimpleNamespace(get=lambda: copy.deepcopy(self.current_cfg)),
                             self.store, CredentialStore(persistent=False), Mock(), demo=True)
        self.engine.reports_dir = self.root / "reports"
        self.app = App(self.engine, self.root, version="test")
        self.outputs = {
            "a": "Synthetic raw A\nuser 60%, system 0%, idle 40%",
            "b": "Synthetic raw B\nuser 60%, system 0%, idle 40%",
        }
        self.names = {"a": "MM <A>", "b": "MM & B"}
        self.rid = self.seed_run()
        self.pruned = []

    def seed_run(self):
        started = time.time() - 3 * 86400
        devices, collections = {}, {}
        for index, did in enumerate(("a", "b"), 1):
            devices[did] = {
                "id": did, "name": self.names[did], "host": f"192.0.2.{index}", "role": "MM",
                "status": "warn", "connected": True, "issues": 1,
                "metrics": {"cpu": 60, "mem": 20,
                            "cpu_detail": {"user": 60, "system": 0, "idle": 40},
                            "mem_detail": {"total_kb": 1000, "used_kb": 200,
                                           "free_kb": 800, "used_pct": 20}},
                "checks": [
                    {"key": "cpu", "label": "CPU", "level": "warn", "summary": "Saved CPU warning",
                     "evidence": [], "command": "show cpu", "cmd_id": "cpu"},
                    {"key": "memory", "label": "Memory", "level": "ok", "summary": "Saved memory normal",
                     "evidence": []},
                ],
            }
            collections[did] = DeviceCollection(
                did, True, started=started, finished=started + 10,
                results={"cpu": CmdResult("cpu", "CPU", "show cpu", self.outputs[did], True)},
            )
        evaluation = {"summary": {"overall": "warn", "counts": {"warn": 2}, "total_devices": 2},
                      "devices": devices}
        return self.store.save_run(started, started + 10, "manual", evaluation, collections, {},
                                   report_config=self.original_cfg)

    def get_response(self, path):
        handler = object.__new__(make_handler(self.app))
        handler.path = path
        handler.headers = {"Host": "localhost:8875"}
        handler._send = Mock()
        handler.do_GET()
        handler._send.assert_called_once()
        return handler._send.call_args.args

    def prune_before_workbook(self, run, site, thresholds, raw):
        # Both old and new implementations already fetched metadata here.
        # Prune before the workbook requests any raw text: the old callback
        # silently returns no output; a captured callback still has both rows.
        self.store.prune(1)
        self.assertIsNone(self.store.get_run(run["id"]))
        self.pruned.append(run["id"])
        return build_workbook(run, site, thresholds, raw)

    def assert_workbook(self, data):
        workbook = load_workbook(io.BytesIO(data))
        try:
            self.assertEqual(workbook.sheetnames, [
                "요약", "백본 CPU·메모리", "컨트롤러 CPU·메모리", "백본 로그", "백본 VRRP", "점검 항목", "원본 출력",
            ])
            self.assertIn(self.original_cfg["site_name"], workbook["요약"]["A1"].value)
            text = "\n".join(str(cell.value) for sheet in workbook for row in sheet for cell in row
                             if cell.value is not None)
            self.assertNotIn(self.current_cfg["site_name"], text)
            self.assertIn("CPU 주의 50% / 장애 70%", text)
            rows = [row for row in workbook["원본 출력"].iter_rows(min_row=5, values_only=True) if row[0]]
            self.assertEqual(len(rows), 2)
            self.assertEqual({row[0]: row[4] for row in rows},
                             {self.names[did]: output for did, output in self.outputs.items()})
            self.assertTrue(all(row[1:4] == ("CPU", "show cpu", "성공") for row in rows))
            self.assertEqual([workbook["컨트롤러 CPU·메모리"].cell(row, 12).value for row in (5, 6)],
                             ["주의", "주의"])
        finally:
            workbook.close()

    def test_xlsx_download_preserves_captured_raw_after_prune(self):
        for explicit_id in (True, False):
            with self.subTest(explicit_id=explicit_id):
                if not explicit_id:
                    self.rid = self.seed_run()
                query = f"?t={self.app.token}" + (f"&id={self.rid}" if explicit_id else "")
                with patch("nc.excel.build_workbook", side_effect=self.prune_before_workbook), \
                        patch.object(self.store, "get_raw", wraps=self.store.get_raw) as live_raw:
                    response = self.get_response("/report.xlsx" + query)
                self.assertEqual(response[0], 200)
                self.assert_workbook(response[1])
                live_raw.assert_not_called()
                self.assertEqual(response[2], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.assertIn("Content-Disposition", response[3])
                self.assertEqual(self.pruned[-1], self.rid)

    def test_excel_save_and_open_preserve_captured_raw_after_prune(self):
        for index, route in enumerate(("/api/excel_save", "/api/excel_open")):
            with self.subTest(route=route):
                if index:
                    self.rid = self.seed_run()
                body = {"id": self.rid} if not index else {}
                with patch("nc.excel.build_workbook", side_effect=self.prune_before_workbook), \
                        patch.object(self.store, "get_raw", wraps=self.store.get_raw) as live_raw, \
                        patch("nc.server._open_path") as open_path:
                    code, payload = self.app.handle_post(route, body)
                self.assertEqual(code, 200)
                self.assertTrue(payload["ok"])
                self.assert_workbook(Path(payload["path"]).read_bytes())
                live_raw.assert_not_called()
                self.assertEqual(open_path.call_count, int(route.endswith("open")))
                self.assertEqual(self.pruned[-1], self.rid)

    def test_automatic_save_preserves_captured_raw_and_notification_after_prune(self):
        cfg = copy.deepcopy(self.current_cfg)
        cfg["report"]["notify"] = True
        expected_notification = completion_message(self.store.get_run(self.rid))
        with patch("nc.excel.build_workbook", side_effect=self.prune_before_workbook), \
                patch.object(self.store, "get_raw", wraps=self.store.get_raw) as live_raw, \
                patch("nc.notify.show_toast") as notify:
            self.engine._after_run(cfg, self.rid)
        self.assertTrue(self.engine.last_report, self.engine.last_error)
        self.assert_workbook(Path(self.engine.last_report).read_bytes())
        live_raw.assert_not_called()
        notify.assert_called_once_with(*expected_notification)
        self.assertEqual(self.engine.last_error, "")
        self.assertEqual(self.pruned, [self.rid])

    def test_html_and_run_api_do_not_fetch_or_decompress_raw_content(self):
        with patch.object(self.store, "get_report_snapshot", side_effect=AssertionError("raw snapshot read")) as snapshot, \
                patch.object(self.store, "get_raw", side_effect=AssertionError("raw read")) as live_raw, \
                patch("nc.store.zlib.decompress", side_effect=AssertionError("raw decompression")) as decompress:
            for suffix in ("", "&download=1"):
                response = self.get_response(f"/report?id={self.rid}&t={self.app.token}" + suffix)
                self.assertEqual(response[0], 200)
                self.assertIn("MM &lt;A&gt;", response[1])
                self.assertIn("MM &amp; B", response[1])
                self.assertIn(self.original_cfg["site_name"], response[1])
                self.assertNotIn(self.current_cfg["site_name"], response[1])
            code, run = self.app.handle_get("/api/run", {"id": str(self.rid)})
        self.assertEqual(code, 200)
        self.assertEqual(set(run["devices"]), {"a", "b"})
        self.assertEqual(run["devices"]["a"]["commands"][0]["cmd_id"], "cpu")
        snapshot.assert_not_called()
        live_raw.assert_not_called()
        decompress.assert_not_called()

    def test_notify_only_does_not_fetch_or_decompress_raw_content(self):
        cfg = copy.deepcopy(self.current_cfg)
        cfg["report"].update(excel_auto=False, notify=True)
        expected_notification = completion_message(self.store.get_run(self.rid))
        with patch.object(self.store, "get_report_snapshot", side_effect=AssertionError("raw snapshot read")) as snapshot, \
                patch.object(self.store, "get_raw", side_effect=AssertionError("raw read")) as live_raw, \
                patch("nc.store.zlib.decompress", side_effect=AssertionError("raw decompression")) as decompress, \
                patch("nc.excel.save_report") as save, patch("nc.notify.show_toast") as notify:
            self.engine._after_run(cfg, self.rid)
        notify.assert_called_once_with(*expected_notification)
        snapshot.assert_not_called()
        live_raw.assert_not_called()
        decompress.assert_not_called()
        save.assert_not_called()
        self.assertEqual(self.engine.last_error, "")

    def test_snapshot_read_failure_remains_an_excel_error_and_keeps_notification(self):
        cfg = copy.deepcopy(self.current_cfg)
        cfg["report"]["notify"] = True
        expected_notification = completion_message(self.store.get_run(self.rid))
        for error in (sqlite3.DatabaseError("synthetic snapshot read failure"), OSError("synthetic I/O failure")):
            with self.subTest(error=type(error).__name__):
                self.engine.last_error = "Earlier synthetic error"
                with patch.object(self.store, "get_report_snapshot", side_effect=error), \
                        patch("nc.excel.save_report") as save, patch("nc.notify.show_toast") as notify:
                    self.engine._after_run(cfg, self.rid)
                save.assert_not_called()
                notify.assert_called_once_with(*expected_notification)
                self.assertTrue(self.engine.last_error.startswith("Earlier synthetic error"))
                self.assertIn("엑셀 보고서를 저장하지 못했습니다", self.engine.last_error)
                self.assertEqual(self.engine.last_report, "")

    def test_corrupt_raw_output_keeps_notification_after_excel_failure(self):
        cfg = copy.deepcopy(self.current_cfg)
        cfg["report"]["notify"] = True
        expected_notification = completion_message(self.store.get_run(self.rid))
        with self.store._db:
            self.store._db.execute("UPDATE raw_outputs SET output=? WHERE run_id=? AND device_id=?",
                                   (b"synthetic invalid compressed output", self.rid, "a"))
        with patch("nc.notify.show_toast") as notify:
            self.engine._after_run(cfg, self.rid)
        notify.assert_called_once_with(*expected_notification)
        self.assertIn("엑셀 보고서를 저장하지 못했습니다", self.engine.last_error)
        self.assertEqual(self.engine.last_report, "")
        self.assertEqual(list(self.engine.reports_dir.glob("*.xlsx")), [])

    def test_prune_before_automatic_snapshot_skips_excel_but_keeps_notification(self):
        cfg = copy.deepcopy(self.current_cfg)
        cfg["report"]["notify"] = True
        expected_notification = completion_message(self.store.get_run(self.rid))
        original_snapshot = self.store.get_report_snapshot

        def prune_then_snapshot(run_id=None):
            self.store.prune(1)
            return original_snapshot(run_id)

        with patch.object(self.store, "get_report_snapshot", side_effect=prune_then_snapshot), \
                patch("nc.excel.save_report") as save, patch("nc.notify.show_toast") as notify:
            self.engine._after_run(cfg, self.rid)
        self.assertIsNone(self.store.get_run(self.rid))
        save.assert_not_called()
        notify.assert_called_once_with(*expected_notification)
        self.assertEqual(self.engine.last_report, "")
        self.assertEqual(self.engine.last_error, "")

    def test_missing_excel_reports_keep_existing_not_found_responses(self):
        self.store.prune(1)
        for query in (f"&id={self.rid}", ""):
            response = self.get_response(f"/report.xlsx?t={self.app.token}" + query)
            self.assertEqual(response[:2], (404, {"error": "점검 기록이 없습니다."}))
        with patch("nc.server._open_path") as open_path, patch("nc.excel.save_report") as save:
            for route in ("/api/excel_save", "/api/excel_open"):
                self.assertEqual(self.app.handle_post(route, {"id": self.rid}),
                                 (404, {"error": "점검 기록이 없습니다."}))
        open_path.assert_not_called()
        save.assert_not_called()


if __name__ == "__main__":
    unittest.main()
