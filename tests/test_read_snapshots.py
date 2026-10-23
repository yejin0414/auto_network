"""Read snapshots survive retention without changing history/report contents.

Only disposable SQLite files and in-memory workbooks are used. Event gates fix
the read/prune ordering; no device, HTTP server, or OS credential store starts.
"""
from __future__ import annotations

import io
import sqlite3
import tempfile
import threading
import time
import unittest
import zlib
from contextlib import ExitStack, closing, contextmanager
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import nc.store as S
from nc.collector import CmdResult, DeviceCollection
from nc.config import validate
from nc.excel import build_workbook


CFG = validate({"site_name": "Synthetic snapshot site", "devices": [
    {"id": "mm", "name": "Synthetic MM", "host": "192.0.2.1", "role": "MM"}
]})


def seed(store, marker="old", *, expired=True):
    stamp = time.time() - (3 * 86400 if expired else 0)
    evaluation = {
        "summary": {"overall": "ok", "counts": {"ok": 1}, "total_devices": 1},
        "devices": {"mm": {
            "id": "mm", "name": "Synthetic MM", "host": "192.0.2.1", "role": "MM",
            "connected": True, "status": "ok", "issues": 0,
            "metrics": {"cpu": 3, "mem": 20}, "checks": [],
        }},
    }
    dc = DeviceCollection("mm", True, results={
        key: CmdResult(key, key, "show " + key, f"SYNTHETIC_{marker}_{key}", True)
        for key in ("version", "cpu")
    })
    return store.save_run(stamp, stamp + 1, "synthetic", evaluation, {"mm": dc}, {}, report_config=CFG)


class _Gate:
    def __init__(self):
        self.entered = threading.Event()
        self.resume = threading.Event()
        self.reader = None

    def pause(self):
        if threading.current_thread() is self.reader and not self.entered.is_set():
            self.entered.set()
            if not self.resume.wait(5):
                raise AssertionError("synthetic read gate was not released")


class _AfterReadUnlock:
    """Prune after the first actual read critical section has been released."""
    def __init__(self, lock, gate):
        self.lock, self.gate = lock, gate

    def acquire(self, *args, **kwargs):
        return self.lock.acquire(*args, **kwargs)

    def release(self):
        self.lock.release()
        self.gate.pause()

    def __enter__(self):
        self.acquire()
        return self

    def __exit__(self, *_):
        self.release()


class _Rows:
    def __init__(self, cursor, after_fetch):
        self.cursor, self.after_fetch = cursor, after_fetch

    def fetchall(self):
        rows = self.cursor.fetchall()
        self.after_fetch()
        return rows

    def fetchone(self):
        row = self.cursor.fetchone()
        self.after_fetch()
        return row

    def __iter__(self):
        return iter(self.fetchall())

    def __getattr__(self, name):
        return getattr(self.cursor, name)


class _Connection:
    """Small hook around a real connection, preserving real SQLite locking."""
    def __init__(self, db, *, before=None, after_parent=None):
        self.db, self.before, self.after_parent = db, before, after_parent

    def execute(self, sql, parameters=()):
        if self.before:
            self.before(sql)
        cursor = self.db.execute(sql, parameters)
        if self.after_parent and sql.lstrip().upper().startswith("SELECT") and "FROM runs" in sql:
            return _Rows(cursor, self.after_parent)
        return cursor

    def __getattr__(self, name):
        return getattr(self.db, name)

    def __enter__(self):
        self.db.__enter__()
        return self

    def __exit__(self, *args):
        return self.db.__exit__(*args)


def workbook_contents(data):
    workbook = load_workbook(io.BytesIO(data))
    try:
        return [(sheet.title, tuple(str(r) for r in sheet.merged_cells.ranges),
                 tuple(tuple((cell.value, cell.data_type, cell.number_format) for cell in row)
                       for row in sheet.iter_rows())) for sheet in workbook]
    finally:
        workbook.close()


class ReadSnapshots(unittest.TestCase):
    @contextmanager
    def stores(self, mode=None, *, separate=False):
        if mode == "WAL" and not S._wal_safe_version(sqlite3.sqlite_version_info):
            self.skipTest(f"SQLite {sqlite3.sqlite_version}: runtime policy uses DELETE; unsupported WAL fixture skipped")
        with ExitStack() as cleanup:
            folder = cleanup.enter_context(tempfile.TemporaryDirectory(prefix="netcheck-read-snapshot-"))
            if mode == "DELETE":
                cleanup.enter_context(patch.object(S, "_wal_safe_version", return_value=False))
            reader = S.Store(Path(folder) / "history.sqlite3")
            cleanup.callback(reader._db.close)
            writer = S.Store(reader.path) if separate else reader
            if separate:
                cleanup.callback(writer._db.close)
            if mode:
                self.assertEqual(reader._db.execute("PRAGMA journal_mode").fetchone()[0], mode.lower())
            yield reader, writer

    def read_during_prune(self, reader, writer, method, run_id):
        load = getattr(reader, method)
        gate, outcome = _Gate(), {}
        separate = reader is not writer
        if separate:
            replacement = _Connection(reader._db, after_parent=gate.pause)
            interception = patch.object(reader, "_db", replacement)
        else:
            interception = patch.object(reader, "_lock", _AfterReadUnlock(reader._lock, gate))
        # A DELETE-mode writer cannot commit while a read snapshot is held.
        # Fail fast, release the reader, and retry instead of deadlocking the gate.
        writer._db.execute("PRAGMA busy_timeout=0")

        def read():
            try:
                outcome["value"] = load(run_id)
            except BaseException as exc:
                outcome["error"] = exc

        gate.reader = threading.Thread(target=read, name="snapshot-test-reader", daemon=True)
        try:
            with interception:
                gate.reader.start()
                try:
                    entered = gate.entered.wait(3)
                    if not entered and "error" in outcome:
                        raise outcome["error"]
                    self.assertTrue(entered, "reader did not reach the controlled boundary")
                    try:
                        writer.prune(1)
                    except sqlite3.OperationalError as exc:
                        self.assertTrue(separate)
                        self.assertEqual(writer._db.execute("PRAGMA journal_mode").fetchone()[0], "delete")
                        self.assertEqual(exc.sqlite_errorcode & 0xff, sqlite3.SQLITE_BUSY)
                        self.assertFalse(writer._db.in_transaction)
                finally:
                    gate.resume.set()
                    gate.reader.join(3)
                self.assertFalse(gate.reader.is_alive(), "snapshot reader did not finish")
            if "error" in outcome:
                raise outcome["error"]
            self.assertFalse(reader._db.in_transaction)
            writer.prune(1)
            self.assertIsNone(reader.get_run(run_id))
            return outcome["value"]
        finally:
            gate.resume.set()
            writer._db.execute("PRAGMA busy_timeout=30000")

    def test_shared_connection_returns_whole_run_after_prune_between_reads(self):
        for mode in ("DELETE", "WAL"):
            with self.subTest(mode=mode), self.stores(mode) as (reader, writer):
                rid = seed(reader)
                expected = reader.get_run(rid)
                actual = self.read_during_prune(reader, writer, "get_run", rid)
                self.assertEqual(actual, expected)

    def test_separate_connection_returns_whole_run_during_prune(self):
        for mode in ("DELETE", "WAL"):
            with self.subTest(mode=mode), self.stores(mode, separate=True) as (reader, writer):
                rid = seed(reader)
                expected = reader.get_run(rid)
                actual = self.read_during_prune(reader, writer, "get_run", rid)
                self.assertEqual(actual, expected)

    def test_report_snapshot_preserves_workbook_after_concurrent_prune(self):
        for mode in ("DELETE", "WAL"):
            for separate in (False, True):
                with self.subTest(mode=mode, separate=separate), self.stores(mode, separate=separate) as (reader, writer):
                    rid = seed(reader)
                    expected_run = reader.get_run(rid)
                    expected_raw = {key: reader.get_raw(rid, "mm", key) for key in ("version", "cpu")}
                    expected_book = build_workbook(expected_run, CFG["site_name"], CFG["thresholds"],
                                                   lambda dev, cmd: expected_raw[cmd])
                    run, raw = self.read_during_prune(reader, writer, "get_report_snapshot", rid)
                    self.assertEqual(run, expected_run)
                    for key, value in expected_raw.items():
                        self.assertEqual(raw("mm", key), value)
                    actual_book = build_workbook(run, CFG["site_name"], CFG["thresholds"], raw)
                    self.assertEqual(workbook_contents(actual_book), workbook_contents(expected_book))

    def test_raw_is_lazy_detached_and_limited_to_the_requested_run(self):
        with self.stores() as (store, _):
            old_id = seed(store, "old")
            new_id = seed(store, "new", expired=False)
            expected = {key: store.get_raw(old_id, "mm", key) for key in ("version", "cpu")}
            # A corrupt BLOB on another run must never be decoded for this report.
            with store._db:
                store._db.execute("UPDATE raw_outputs SET output=? WHERE run_id=?", (b"synthetic-invalid-zlib", new_id))
            with patch("nc.store.zlib.decompress", wraps=zlib.decompress) as decompress:
                run, raw = store.get_report_snapshot(old_id)
                decompress.assert_not_called()
                self.assertEqual(run["id"], old_id)
                store._db.close()
                for key, value in expected.items():
                    self.assertEqual(raw("mm", key), value)
                self.assertEqual(decompress.call_count, 2)
                self.assertIsNone(raw("missing-device", "version"))
                self.assertIsNone(raw("mm", "missing-command"))
                self.assertEqual(decompress.call_count, 2)

    def test_normal_history_does_not_select_or_decompress_output_blobs(self):
        with self.stores() as (store, _):
            rid = seed(store)
            expected = store.get_run(rid)
            forbidden_reads = []

            def authorize(action, table, column, *_):
                if action == sqlite3.SQLITE_READ and table == "raw_outputs" and column == "output":
                    forbidden_reads.append((table, column))
                    return sqlite3.SQLITE_DENY
                return sqlite3.SQLITE_OK

            store._db.set_authorizer(authorize)
            try:
                with patch("nc.store.zlib.decompress", side_effect=AssertionError("ordinary history decoded a BLOB")):
                    self.assertEqual(store.get_run(rid), expected)
                    self.assertEqual(store.list_runs()[0]["id"], rid)
            finally:
                store._db.set_authorizer(None)
            self.assertEqual(forbidden_reads, [])

    def assert_reusable(self, store, rid, expected, *, outer=False):
        self.assertEqual(store._db.in_transaction, outer)
        self.assertTrue(store._lock.acquire(timeout=1), "read failure retained the Store lock")
        store._lock.release()
        self.assertEqual(store.get_run(rid), expected)
        self.assertEqual(store._db.in_transaction, outer)

    def test_query_failure_releases_snapshot_and_preserves_original_error(self):
        for method in ("get_run", "get_report_snapshot"):
            with self.subTest(method=method), self.stores() as (store, _):
                rid = seed(store)
                expected = store.get_run(rid)
                error = sqlite3.OperationalError("synthetic snapshot query failure")
                failed = []

                def fail_once(sql):
                    if "FROM device_results" in sql and not failed:
                        failed.append(sql)
                        raise error

                with patch.object(store, "_db", _Connection(store._db, before=fail_once)):
                    with self.assertRaises(sqlite3.OperationalError) as caught:
                        getattr(store, method)(rid)
                    self.assertIs(caught.exception, error)
                    self.assert_reusable(store, rid, expected)
                self.assertIsInstance(seed(store, "recovered", expired=False), int)

    def test_release_failures_leave_connection_reusable_without_outer_transaction(self):
        for failures in (1, 2):
            with self.subTest(failures=failures), self.stores() as (store, _):
                rid = seed(store)
                expected = store.get_run(rid)
                errors = [sqlite3.OperationalError(f"synthetic release failure {n}") for n in range(failures)]
                original = errors[0]

                def fail_release(sql):
                    if sql.lstrip().upper().startswith("RELEASE") and errors:
                        raise errors.pop(0)

                with patch.object(store, "_db", _Connection(store._db, before=fail_release)):
                    with self.assertRaises(sqlite3.OperationalError) as caught:
                        store.get_report_snapshot(rid)
                    self.assertIs(caught.exception, original)
                    self.assertEqual(errors, [])
                    self.assert_reusable(store, rid, expected)
                self.assertIsInstance(seed(store, "recovered", expired=False), int)

    def test_outer_transaction_is_neither_committed_nor_rolled_back_by_snapshot(self):
        for failure in (None, "query", "release"):
            with self.subTest(failure=failure), self.stores("DELETE") as (store, _):
                rid = seed(store)
                expected = store.get_run(rid)
                db = store._db
                db.execute("BEGIN")
                db.execute("INSERT INTO events(ts,text) VALUES(?,?)", (time.time(), "uncommitted outer change"))
                remaining = [2 if failure == "release" else 1]
                error = sqlite3.OperationalError("synthetic inner snapshot failure")

                def fail(sql):
                    selected = ((failure == "query" and "FROM device_results" in sql)
                                or (failure == "release" and sql.lstrip().upper().startswith("RELEASE")))
                    if selected and remaining[0]:
                        remaining[0] -= 1
                        raise error

                try:
                    with patch.object(store, "_db", _Connection(db, before=fail)):
                        if failure:
                            with self.assertRaises(sqlite3.OperationalError) as caught:
                                store.get_report_snapshot(rid)
                            self.assertIs(caught.exception, error)
                        else:
                            self.assertEqual(store.get_report_snapshot(rid)[0], expected)
                        self.assert_reusable(store, rid, expected, outer=True)
                        self.assertEqual(db.execute("SELECT COUNT(*) FROM events").fetchone()[0], 1)
                    with closing(sqlite3.connect(store.path)) as other:
                        self.assertEqual(other.execute("SELECT COUNT(*) FROM events").fetchone()[0], 0)
                finally:
                    db.rollback()
                self.assertFalse(db.in_transaction)
                self.assertEqual(db.execute("SELECT COUNT(*) FROM events").fetchone()[0], 0)
                self.assertEqual(store.get_run(rid), expected)

    def test_latest_and_missing_run_keep_existing_return_shapes(self):
        with self.stores() as (store, _):
            self.assertIsNone(store.get_run())
            empty, raw = store.get_report_snapshot()
            self.assertIsNone(empty)
            self.assertIsNone(raw("mm", "version"))
            seed(store, "old")
            latest = seed(store, "latest", expired=False)
            self.assertEqual(store.get_run(), store.get_run(latest))
            run, raw = store.get_report_snapshot()
            self.assertEqual(run, store.get_run(latest))
            self.assertEqual(raw("mm", "version"), store.get_raw(latest, "mm", "version"))
            self.assertIsNone(store.get_run(latest + 100))
            missing, raw = store.get_report_snapshot(latest + 100)
            self.assertIsNone(missing)
            self.assertIsNone(raw("mm", "version"))
            self.assertFalse(store._db.in_transaction)


if __name__ == "__main__":
    unittest.main()
