"""Unknown optional observations survive persistence and exported reports."""
import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from nc.collector import CmdResult, DeviceCollection
from nc.commands import BB_COMMANDS
from nc.config import validate
from nc.evaluate import evaluate_run
from nc.excel import build_workbook
from nc.report import render
from nc.store import Store


class OptionalHistoryReports(unittest.TestCase):
    def test_unknown_observations_preserve_baseline_raw_and_exported_results(self):
        keys = ("device", "power", "fan", "lagg", "ospf")
        specs = {s.id: s for s in BB_COMMANDS}
        cfg = validate({"devices": [{"id": "bb", "name": "Synthetic BB", "role": "BB",
                                     "host": "192.0.2.10"}]})
        baseline = {"bb": {"slots": ["1", "2"], "lagg": {"BAGG1": 2},
                           "power_states": {"1/1": "Normal"}, "fan_states": {"1/1": "Normal"},
                           "ospf": ["192.0.2.1"]}}
        for command_ok in (True, False):
            with self.subTest(command_ok=command_ok), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "history.sqlite3"
                store = Store(path)
                try:
                    store.save_baselines(baseline)
                    dc = DeviceCollection("bb", True, started=1700000000, finished=1700000010)
                    for key in keys:
                        spec = specs[key]
                        dc.results[key] = CmdResult(key, spec.label, spec.commands[0],
                                                    f"unparsed synthetic {key} output", command_ok,
                                                    "" if command_ok else "synthetic command failure", optional=True)
                    evaluation, updates = evaluate_run(cfg, {"bb": dc}, {}, store.baselines(), {},
                                                       now=1700000010)
                    store.save_baselines(updates)
                    run_id = store.save_run(1700000000, 1700000010, "test", evaluation, {"bb": dc}, {},
                                            report_config=cfg)
                finally:
                    store._db.close()

                reopened = Store(path)
                try:
                    run = reopened.get_run(run_id)
                    self.assertEqual(reopened.baselines(), baseline)
                    self.assertEqual(run["overall"], "unknown")
                    checks = {c["key"]: c for c in run["devices"]["bb"]["checks"]}
                    self.assertEqual(set(checks), set(keys))
                    for key in keys:
                        self.assertEqual(checks[key]["level"], "unknown")
                        self.assertEqual(checks[key]["cmd_id"], key)
                        self.assertEqual(reopened.get_raw(run_id, "bb", key)["output"], dc.results[key].output)

                    html = render(run, "Current site")
                    self.assertGreaterEqual(html.count("확인 불가"), len(keys))
                    for key in keys:
                        self.assertIn(checks[key]["label"], html)
                    workbook = load_workbook(io.BytesIO(build_workbook(
                        run, "Current site", cfg["thresholds"],
                        lambda device, command: reopened.get_raw(run_id, device, command))))
                    rows = {row[2]: row[3] for row in workbook["점검 항목"].iter_rows(min_row=5, values_only=True)
                            if row[2]}
                    self.assertEqual(rows, {checks[key]["label"]: "확인 불가" for key in keys})
                    self.assertEqual(workbook["요약"]["D5"].value, "확인 불가")
                finally:
                    reopened._db.close()


if __name__ == "__main__":
    unittest.main()
